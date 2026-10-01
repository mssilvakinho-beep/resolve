from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font

def build_excel(summary, transactions, cases):
    wb=Workbook(); ws=wb.active; ws.title='RESUMO'
    for cell in ws[1]: cell.font=Font(bold=True)
    rows=[('Indicador','Valor'),('Receitas',summary['receitas']),('Despesas',summary['despesas']),('Saldo registrado',summary['saldo_registrado']),('A receber',summary['a_receber']),('A pagar',summary['a_pagar'])]
    for r in rows: ws.append(r)
    tx=wb.create_sheet('MOVIMENTAÇÕES'); tx.append(['ID','Tipo','Categoria','Descrição','Valor','Caso','Vencimento','Pago','Criado em'])
    for c in tx[1]: c.font=Font(bold=True)
    for x in transactions: tx.append([x['id'],x['kind'],x['category'],x['description'],x['amount'],x['case_id'],x['due_date'],bool(x['paid']),x['created_at']])
    cs=wb.create_sheet('CASOS'); cs.append(['ID','Título','Status','Descrição','Criado em','Atualizado em'])
    for c in cs[1]: c.font=Font(bold=True)
    for x in cases: cs.append([x['id'],x['title'],x['status'],x['description'],x['created_at'],x['updated_at']])
    for sheet in wb.worksheets:
        for col in sheet.columns:
            letter=col[0].column_letter; sheet.column_dimensions[letter].width=min(max(len(str(c.value or '')) for c in col)+2,50)
    out=BytesIO(); wb.save(out); return out.getvalue()
